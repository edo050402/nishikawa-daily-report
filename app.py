from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
import sqlite3, os, io, shutil
from datetime import datetime, date, timedelta
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("DAILY_REPORT_DB", os.path.join(BASE_DIR, "daily_report.db"))
SECRET_KEY = os.environ.get("SECRET_KEY", "change-this-secret-key-before-public-deployment")

app = Flask(__name__)
app.secret_key = SECRET_KEY


EXCEL_DIR = os.path.join(BASE_DIR, "excel")
EXCEL_BACKUP_DIR = os.path.join(EXCEL_DIR, "backup")
FULLWIDTH_MONTH_SHEETS = {1:"１月", 2:"２月", 3:"３月", 4:"４月", 5:"５月", 6:"６月", 7:"７月", 8:"８月", 9:"９月", 10:"１０月", 11:"１１月", 12:"１２月"}

SCHEDULE_DIR = os.path.join(BASE_DIR, "schedule_excel")
WORK_SCHEDULE_FILE = os.environ.get("WORK_SCHEDULE_FILE", os.path.join(SCHEDULE_DIR, "日程.xls"))
VEHICLE_SCHEDULE_FILE = os.environ.get("VEHICLE_SCHEDULE_FILE", os.path.join(SCHEDULE_DIR, "車輛日程.xlsx"))



def fiscal_year_for(work_date):
    d = datetime.strptime(work_date, "%Y-%m-%d").date()
    return d.year if d.month >= 5 else d.year - 1


_EXCEL_EMPLOYEE_CACHE = {}

def _excel_employee_no(path):
    """Read the employee number cached in システム!C4 without changing the xlsm."""
    stat = os.stat(path)
    key = (os.path.abspath(path), stat.st_mtime_ns, stat.st_size)
    if key in _EXCEL_EMPLOYEE_CACHE:
        return _EXCEL_EMPLOYEE_CACHE[key]
    try:
        wb = load_workbook(path, read_only=True, data_only=True, keep_vba=True)
        try:
            if "システム" not in wb.sheetnames:
                value = None
            else:
                value = wb["システム"]["C4"].value
        finally:
            wb.close()
        value = "" if value is None else str(value).strip()
    except Exception:
        value = ""
    _EXCEL_EMPLOYEE_CACHE[key] = value
    return value

def find_excel_file(work_date, employee_no=None):
    if not os.path.isdir(EXCEL_DIR):
        raise FileNotFoundError("excelフォルダーがありません。")
    fiscal_year = str(fiscal_year_for(work_date))
    files = [f for f in os.listdir(EXCEL_DIR) if f.lower().endswith(".xlsm") and not f.startswith("~$")]
    preferred = [f for f in files if fiscal_year in f]
    candidates = preferred or files
    if not candidates:
        raise FileNotFoundError("excelフォルダーに業務日報(.xlsm)がありません。")

    # V3.9.6: generated workbooks carry 社員番号 in the filename, so they can be
    # resolved instantly without opening every .xlsm just to inspect システム!C4.
    # Older workbooks are still supported via the C4 fallback below.
    if employee_no is not None and str(employee_no).strip():
        wanted = str(employee_no).strip()
        filename_token = f"_社員番号{wanted}"
        filename_matches = [
            os.path.abspath(os.path.join(EXCEL_DIR, f))
            for f in candidates
            if filename_token in os.path.splitext(f)[0]
        ]
        if filename_matches:
            filename_matches.sort(key=os.path.getmtime, reverse=True)
            return filename_matches[0]

        matched = []
        for filename in candidates:
            path = os.path.abspath(os.path.join(EXCEL_DIR, filename))
            if _excel_employee_no(path) == wanted:
                matched.append(path)
        if not matched:
            raise FileNotFoundError(f"社員番号 {wanted} の業務日報Excelがexcelフォルダーに見つかりません。")
        if len(matched) > 1:
            matched.sort(key=os.path.getmtime, reverse=True)
        return matched[0]

    if len(candidates) > 1:
        candidates.sort(key=lambda f: os.path.getmtime(os.path.join(EXCEL_DIR, f)), reverse=True)
    return os.path.abspath(os.path.join(EXCEL_DIR, candidates[0]))



def _safe_excel_name(value):
    value = str(value or "").strip()
    for ch in '<>:"/\\|?*':
        value = value.replace(ch, "_")
    return value or "employee"

def _employee_excel_template(fiscal_year):
    template_dir = os.path.join(EXCEL_DIR, "template")
    path = os.path.join(template_dir, f"{fiscal_year}_業務日報_template.xlsm")
    if not os.path.isfile(path):
        raise FileNotFoundError(
            f"{fiscal_year}年度のExcelテンプレートがありません：excel\\template\\{fiscal_year}_業務日報_template.xlsm"
        )
    return os.path.abspath(path)

def create_employee_excel(employee_no, name, fiscal_year=None):
    """Create a per-employee workbook by direct file copy only.

    V3.9.7 intentionally does NOT launch Microsoft Excel and does NOT clean the
    workbook. This makes account creation and bulk creation fast and avoids COM /
    merged-cell lockups. The current template contents are copied as-is, as requested.
    The employee number is embedded in the filename and is used for safe matching.
    """
    employee_no = str(employee_no or "").strip()
    name = str(name or "").strip()
    if not employee_no:
        raise ValueError("社員番号を入力してください。")
    if not name:
        raise ValueError("氏名を入力してください。")

    fy = int(fiscal_year or fiscal_year_for(date.today().isoformat()))
    os.makedirs(EXCEL_DIR, exist_ok=True)

    # V3.9.7: bulk creation must NEVER open/parse existing .xlsm files.
    # The previous version called find_excel_file() here; its compatibility fallback
    # parsed workbooks with openpyxl and made bulk creation take minutes.
    token = f"_社員番号{employee_no}"
    for existing_name in os.listdir(EXCEL_DIR):
        if (existing_name.lower().endswith(".xlsm")
                and not existing_name.startswith("~$")
                and str(fy) in existing_name
                and token in os.path.splitext(existing_name)[0]):
            return os.path.abspath(os.path.join(EXCEL_DIR, existing_name)), False

    template = _employee_excel_template(fy)
    filename = f"{fy}_業務日報({_safe_excel_name(name)})_社員番号{employee_no}.xlsm"
    dest = os.path.abspath(os.path.join(EXCEL_DIR, filename))

    # Pure filesystem copy only: no Excel COM and no openpyxl workbook parsing.
    shutil.copy2(template, dest)
    return dest, True

def excel_time_value(hhmm):
    if not hhmm:
        return None
    t = datetime.strptime(hhmm, "%H:%M")
    return (t.hour * 60 + t.minute) / 1440.0


def sync_report_to_company_excel(rid):
    # Uses the installed Microsoft Excel application itself (Windows COM),
    # so the original .xlsm macros, formulas, formatting and validations stay intact.
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        raise RuntimeError("Excel連携にはWindows版Microsoft Excelとpywin32が必要です。requirements.txtを再インストールしてください。")

    conn = db()
    report = conn.execute("""SELECT d.*, u.employee_no, u.name
                             FROM daily_reports d JOIN users u ON u.id=d.user_id
                             WHERE d.id=?""", (rid,)).fetchone()
    entries = conn.execute("""SELECT e.*, j.project_name, j.client_name
                              FROM work_entries e LEFT JOIN jobs j ON j.id=e.job_id
                              WHERE e.daily_report_id=? ORDER BY e.sort_order,e.id""", (rid,)).fetchall()
    adjustments = conn.execute("""SELECT * FROM time_adjustments
                                  WHERE daily_report_id=? ORDER BY sort_order,id""", (rid,)).fetchall()
    conn.close()
    if not report:
        raise ValueError("日報データが見つかりません。")
    if len(entries) > 7:
        raise ValueError("会社Excelは1日7件までの作業欄です。8件以上はデータベースには保存されましたがExcelへ書き込めません。")
    if len(adjustments) > 14:
        raise ValueError("会社Excelは1日最大14件（各作業2件）の時間区分まで反映できます。")

    excel_path = find_excel_file(report["work_date"], report["employee_no"])
    d = datetime.strptime(report["work_date"], "%Y-%m-%d").date()
    sheet_name = FULLWIDTH_MONTH_SHEETS[d.month]

    # Keep a recoverable copy before every automatic write.
    os.makedirs(EXCEL_BACKUP_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_name = f"{os.path.splitext(os.path.basename(excel_path))[0]}_{stamp}.xlsm"
    shutil.copy2(excel_path, os.path.join(EXCEL_BACKUP_DIR, backup_name))

    pythoncom.CoInitialize()
    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(excel_path, UpdateLinks=0, ReadOnly=False)

        # V3.9.6: generated files are copied instantly. When a report is actually
        # written, Excel is already open, so update the owner identity at no extra
        # startup cost. This also converts a copied template into the employee's file.
        try:
            sys_ws = workbook.Worksheets("システム")
            sys_ws.Range("B3").Value = os.path.basename(excel_path)
            sys_ws.Range("B4").Value = report["name"] or ""
            sys_ws.Range("C4").Value = str(report["employee_no"] or "")
        except Exception:
            pass

        try:
            ws = workbook.Worksheets(sheet_name)
        except Exception:
            raise ValueError(f"Excelに{sheet_name}シートがありません。")

        # Day 1 starts at row 7. Every calendar day occupies 22 rows.
        base = 7 + (d.day - 1) * 22
        ws.Cells(base + 1, 3).Value = excel_time_value(report["attendance_start"])
        ws.Cells(base + 1, 5).Value = excel_time_value(report["attendance_end"])
        ws.Cells(base + 1, 3).NumberFormat = "hh:mm"
        ws.Cells(base + 1, 5).NumberFormat = "hh:mm"

        # Clear only the input cells. Existing formulas and formatting are not touched.
        for slot in range(7):
            top = base + 5 + slot * 2
            time_row = top + 1
            for col in (2, 3, 6):
                ws.Cells(top, col).Value = None
            ws.Cells(time_row, 3).Value = None
            ws.Cells(time_row, 5).Value = None
            ws.Cells(top, 9).Value = "-"
            ws.Cells(top, 10).Value = None
            ws.Cells(time_row, 9).Value = "-"
            ws.Cells(time_row, 10).Value = None

        # B in the original sheet is "JOB番号 or 業務コード".
        for slot, entry in enumerate(entries):
            top = base + 5 + slot * 2
            time_row = top + 1
            identifier = entry["job_no_snapshot"] or entry["work_code"] or ""
            ws.Cells(top, 2).Value = identifier
            ws.Cells(top, 3).Value = entry["content"] or ""
            ws.Cells(top, 6).Value = entry["location"] or ""
            ws.Cells(time_row, 3).Value = excel_time_value(entry["start_time"])
            ws.Cells(time_row, 5).Value = excel_time_value(entry["end_time"])
            ws.Cells(time_row, 3).NumberFormat = "hh:mm"
            ws.Cells(time_row, 5).NumberFormat = "hh:mm"

        # 時間区分 belongs to each 作業明細. Excel has two I/J rows per work slot,
        # so each work item can carry up to two manual classifications (e.g. み + 残).
        grouped = {}
        for adj in adjustments:
            grouped.setdefault(int(adj["work_entry_order"] or 0), []).append(adj)
        for entry_order, items in grouped.items():
            if entry_order < 0 or entry_order >= 7:
                continue
            if len(items) > 2:
                raise ValueError(f"作業{entry_order+1}の時間区分はExcel上2件までです。")
            top = base + 5 + entry_order * 2
            for pos, adj in enumerate(items):
                row = top + pos
                typ = adj["category"] or "-"
                ws.Cells(row, 9).Value = typ
                ws.Cells(row, 10).Value = float(adj["hours"] or 0) if typ != "-" else None
                if typ != "-":
                    ws.Cells(row, 10).NumberFormat = "0.00"

        workbook.Save()
        conn2=db(); conn2.execute("UPDATE daily_reports SET sync_status='synced' WHERE id=?",(rid,)); conn2.commit(); conn2.close()
        return excel_path
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()

WORK_CODES = {
    "A": "見積り",
    "B": "客先打合せ",
    "C": "社内会議・打合せ",
    "D": "クレーム対応",
    "E": "移動",
    "F": "時間外移動",
    "G": "その他",
    "H1": "午前有給",
    "H2": "午後有給",
}


def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def table_exists(conn, name):
    return conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (name,)).fetchone() is not None


def ensure_column(conn, table, column, definition):
    cols = {row["name"] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}
    if column not in cols:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_db():
    conn = db()
    conn.executescript('''
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        employee_no TEXT UNIQUE,
        name TEXT NOT NULL,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        role TEXT NOT NULL DEFAULT 'employee',
        active INTEGER NOT NULL DEFAULT 1,
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS jobs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        job_no TEXT UNIQUE NOT NULL,
        project_name TEXT,
        client_name TEXT,
        default_content TEXT,
        active INTEGER NOT NULL DEFAULT 1,
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS daily_reports (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER NOT NULL,
        work_date TEXT NOT NULL,
        attendance_start TEXT NOT NULL DEFAULT '08:30',
        attendance_end TEXT NOT NULL DEFAULT '17:30',
        remarks TEXT,
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
        updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(user_id, work_date),
        FOREIGN KEY(user_id) REFERENCES users(id)
    );

    CREATE TABLE IF NOT EXISTS work_entries (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        daily_report_id INTEGER NOT NULL,
        job_id INTEGER,
        job_no_snapshot TEXT,
        work_code TEXT NOT NULL DEFAULT '',
        content TEXT NOT NULL,
        location TEXT NOT NULL DEFAULT '',
        start_time TEXT,
        end_time TEXT,
        work_minutes INTEGER NOT NULL DEFAULT 0,
        sort_order INTEGER NOT NULL DEFAULT 0,
        FOREIGN KEY(daily_report_id) REFERENCES daily_reports(id) ON DELETE CASCADE,
        FOREIGN KEY(job_id) REFERENCES jobs(id)
    );
    CREATE INDEX IF NOT EXISTS idx_daily_user_date ON daily_reports(user_id, work_date);
    CREATE INDEX IF NOT EXISTS idx_entries_report ON work_entries(daily_report_id, sort_order);
    CREATE TABLE IF NOT EXISTS time_adjustments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        daily_report_id INTEGER NOT NULL,
        category TEXT NOT NULL DEFAULT '-',
        hours REAL NOT NULL DEFAULT 0,
        sort_order INTEGER NOT NULL DEFAULT 0,
        work_entry_order INTEGER NOT NULL DEFAULT 0,
        FOREIGN KEY(daily_report_id) REFERENCES daily_reports(id) ON DELETE CASCADE
    );
    CREATE INDEX IF NOT EXISTS idx_adjust_report ON time_adjustments(daily_report_id, sort_order);

    CREATE TABLE IF NOT EXISTS schedule_work_cache (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        schedule_date TEXT NOT NULL,
        person_name TEXT NOT NULL,
        item_text TEXT NOT NULL,
        sort_order INTEGER NOT NULL DEFAULT 0
    );
    CREATE INDEX IF NOT EXISTS idx_schedule_work_date ON schedule_work_cache(schedule_date, sort_order);

    CREATE TABLE IF NOT EXISTS schedule_vehicle_cache (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        schedule_date TEXT NOT NULL,
        vehicle TEXT NOT NULL,
        plate TEXT,
        default_user TEXT,
        assigned TEXT,
        sort_order INTEGER NOT NULL DEFAULT 0
    );
    CREATE INDEX IF NOT EXISTS idx_schedule_vehicle_date ON schedule_vehicle_cache(schedule_date, sort_order);

    CREATE TABLE IF NOT EXISTS schedule_sync_state (
        id INTEGER PRIMARY KEY CHECK (id=1),
        synced_at TEXT,
        work_file_mtime REAL,
        vehicle_file_mtime REAL
    );
    ''')

    ensure_column(conn, "work_entries", "location", "TEXT NOT NULL DEFAULT ''")
    ensure_column(conn, "time_adjustments", "work_entry_order", "INTEGER NOT NULL DEFAULT 0")
    ensure_column(conn, "users", "theme", "TEXT NOT NULL DEFAULT 'light'")
    ensure_column(conn, "users", "language", "TEXT NOT NULL DEFAULT 'ja'")
    ensure_column(conn, "daily_reports", "sync_status", "TEXT NOT NULL DEFAULT 'unsynced'")

    # One-time migration of V1 report rows, if a V1 database is reused.
    if table_exists(conn, "reports"):
        migrated = conn.execute("SELECT COUNT(*) AS n FROM work_entries").fetchone()["n"]
        if migrated == 0:
            rows = conn.execute("SELECT * FROM reports ORDER BY user_id, work_date, start_time").fetchall()
            for r in rows:
                dr = conn.execute("SELECT id FROM daily_reports WHERE user_id=? AND work_date=?", (r["user_id"], r["work_date"])).fetchone()
                if not dr:
                    conn.execute("INSERT INTO daily_reports(user_id,work_date,attendance_start,attendance_end,remarks) VALUES(?,?,?,?,?)",
                                 (r["user_id"], r["work_date"], r["start_time"] or "08:30", r["end_time"] or "17:30", r["remarks"] or ""))
                    dr_id = conn.execute("SELECT last_insert_rowid() AS id").fetchone()["id"]
                else:
                    dr_id = dr["id"]
                conn.execute('''INSERT INTO work_entries(daily_report_id,job_no_snapshot,work_code,content,start_time,end_time,work_minutes,sort_order)
                                VALUES(?,?,?,?,?,?,?,?)''',
                             (dr_id, r["job_no"], r["work_code"], r["content"], r["start_time"], r["end_time"], r["work_minutes"], 0))
    conn.commit(); conn.close()


def has_users():
    conn = db(); row = conn.execute("SELECT COUNT(*) AS n FROM users").fetchone(); conn.close(); return row["n"] > 0


def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapper


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        if session.get("role") != "admin":
            flash("管理者のみ利用できます。", "error")
            return redirect(url_for("dashboard"))
        return fn(*args, **kwargs)
    return wrapper


def parse_hhmm(value):
    return datetime.strptime(value, "%H:%M") if value else None


def valid_quarter(value):
    if not value: return True
    try: return parse_hhmm(value).minute in (0, 15, 30, 45)
    except ValueError: return False


def to_minutes(value):
    t = parse_hhmm(value)
    return t.hour * 60 + t.minute if t else None


def minutes_between(start, end):
    s = to_minutes(start); e = to_minutes(end)
    if s is None or e is None: return 0
    if e <= s: e += 1440
    return e - s


def interval(start, end):
    s = to_minutes(start); e = to_minutes(end)
    if s is None or e is None: return None
    if e <= s: e += 1440
    return s, e


def intervals_overlap(a, b):
    return max(a[0], b[0]) < min(a[1], b[1])




UI_TEXT = {
    "ja": {
        "daily_list":"日報一覧","job_no":"JOB番号","schedule":"スケジュール","settings":"設定","monthly_status":"月次勤務状況","month":"月","show":"表示",
        "work_days":"勤務日","day":"日","work_hours":"作業時間","hours":"時間","time_summary":"時間区分合計","deemed":"みなし","overtime":"残業","early":"早出","f_move":"F 時間外移動","night":"深夜",
        "entered_reports":"入力済み日報","excel_save_all":"Excelへ一括保存","excel_import":"Excelから取込","synced":"同期済み","imported":"Excel取込","unsynced":"未同期",
        "edit":"編集","copy":"コピー","delete":"削除","no_reports":"この月の日報はまだありません。","daily_entry":"日報入力",
        "display_settings":"表示設定","theme":"テーマ","light":"ライト","dark":"ダーク","language":"言語","save_settings":"設定を保存",
        "management":"管理","account_management":"アカウント管理","job_management":"JOB番号管理","db_backup":"データベースバックアップ","session":"セッション","logout":"ログアウト",
        "settings_desc":"表示・言語・アカウントを管理します。","account_desc":"利用者アカウントを作成・編集できます。","new_account":"＋ 新規アカウント",
        "employee_no":"社員番号","name":"氏名","username":"ユーザー名","password":"パスワード","role":"権限","employee":"一般","admin":"管理者","create":"作成","registered_accounts":"登録済みアカウント","actions":"操作","excel_accounts":"社員別Excel","create_missing_excel":"不足Excelを一括作成","excel_auto_note":"新規アカウント作成時、社員番号に対応する業務日報Excelも自動作成します。",
        "job_db":"JOB番号データベース","job_desc":"登録したJOB番号は日報入力時に一覧から選択できます。","add_job":"＋ JOB追加","project":"案件名","client":"客先名","default_content":"標準内容","register":"登録","registered_jobs":"登録済みJOB","status":"状態","active":"有効","inactive":"無効",
        "job_edit":"JOB編集","back":"戻る","save":"保存","account_edit":"アカウント編集","new_password":"新しいパスワード（変更しない場合は空欄）",
        "report_edit":"日報編集","report_input":"業務日報入力","report_help":"1日に複数のJOB・時間帯を入力できます。時刻・時間は15分単位で入力してください。","quarter":"15分単位","date":"日付","start_work":"勤務開始","end_work":"勤務終了","work_details":"作業明細","add_work":"＋ 作業を追加","remarks":"備考",
        "no_job":"JOBなし","work_code":"業務コード（任意）","no_select":"選択しない","content":"内容","place":"場所","choose":"選択してください","office":"社内","site":"現場","start":"開始","end":"終了","time_class":"時間区分","add_class":"＋ 区分を追加","class_help":"この作業に対する区分を手動で選択します。","classification":"区分","login":"ログイン","report_login":"業務日報 ログイン",
        "schedule_title":"スケジュール","schedule_desc":"社内作業予定と車両予定をExcelから表示します。","schedule_date":"表示日","today":"今日","work_schedule":"社内作業予定","vehicle_schedule":"車両予定","scheduled_people":"予定あり","people":"名","no_work_schedule":"この日の社内予定はありません。","vehicle":"車両","plate":"車両番号","user":"使用者","vehicle_status":"状態","available":"空き","in_use_plan":"使用予定","source_file":"参照Excel","excel_read_error":"Excelを読み込めませんでした","schedule_refresh":"Excelから更新","schedule_refreshing":"Excelからデータベースへ更新します","schedule_last_sync":"最終更新","schedule_not_synced":"未更新","schedule_sync_done":"スケジュールをデータベースへ更新しました。"
    },
    "id": {
        "daily_list":"Daftar Laporan","job_no":"Nomor JOB","schedule":"Jadwal","settings":"Pengaturan","monthly_status":"Status Kerja Bulanan","month":"Bulan","show":"Tampilkan",
        "work_days":"Hari Kerja","day":"hari","work_hours":"Jam Kerja","hours":"jam","time_summary":"Total Kategori Waktu","deemed":"Minashi","overtime":"Lembur","early":"Masuk Awal","f_move":"F Perjalanan di Luar Jam","night":"Malam",
        "entered_reports":"Laporan yang Sudah Diisi","excel_save_all":"Simpan Semua ke Excel","excel_import":"Impor dari Excel","synced":"Sudah Sinkron","imported":"Impor Excel","unsynced":"Belum Sinkron",
        "edit":"Edit","copy":"Salin","delete":"Hapus","no_reports":"Belum ada laporan untuk bulan ini.","daily_entry":"Input Laporan",
        "display_settings":"Pengaturan Tampilan","theme":"Tema","light":"Terang","dark":"Gelap","language":"Bahasa","save_settings":"Simpan Pengaturan",
        "management":"Manajemen","account_management":"Manajemen Akun","job_management":"Manajemen Nomor JOB","db_backup":"Backup Database","session":"Sesi","logout":"Keluar",
        "settings_desc":"Kelola tampilan, bahasa, dan akun.","account_desc":"Buat dan edit akun pengguna.","new_account":"＋ Akun Baru",
        "employee_no":"Nomor Karyawan","name":"Nama","username":"Nama Pengguna","password":"Kata Sandi","role":"Hak Akses","employee":"Umum","admin":"Admin","create":"Buat","registered_accounts":"Akun Terdaftar","actions":"Aksi","excel_accounts":"Excel per Karyawan","create_missing_excel":"Buat Excel yang Belum Ada","excel_auto_note":"Saat membuat akun baru, Excel laporan harian sesuai Nomor Karyawan juga dibuat otomatis.",
        "job_db":"Database Nomor JOB","job_desc":"Nomor JOB yang terdaftar dapat dipilih saat mengisi laporan.","add_job":"＋ Tambah JOB","project":"Nama Proyek","client":"Nama Pelanggan","default_content":"Isi Standar","register":"Daftar","registered_jobs":"JOB Terdaftar","status":"Status","active":"Aktif","inactive":"Nonaktif",
        "job_edit":"Edit JOB","back":"Kembali","save":"Simpan","account_edit":"Edit Akun","new_password":"Kata sandi baru (kosongkan jika tidak diubah)",
        "report_edit":"Edit Laporan","report_input":"Input Laporan Kerja","report_help":"Dalam satu hari dapat diisi beberapa JOB dan rentang waktu. Masukkan waktu dalam satuan 15 menit.","quarter":"Satuan 15 menit","date":"Tanggal","start_work":"Mulai Kerja","end_work":"Selesai Kerja","work_details":"Rincian Pekerjaan","add_work":"＋ Tambah Pekerjaan","remarks":"Catatan",
        "no_job":"Tanpa JOB","work_code":"Kode Kerja (opsional)","no_select":"Tidak dipilih","content":"Isi","place":"Lokasi","choose":"Pilih","office":"Kantor","site":"Lapangan","start":"Mulai","end":"Selesai","time_class":"Kategori Waktu","add_class":"＋ Tambah Kategori","class_help":"Pilih kategori untuk pekerjaan ini secara manual.","classification":"Kategori","login":"Masuk","report_login":"Login Laporan Kerja",
        "schedule_title":"Jadwal","schedule_desc":"Menampilkan jadwal kerja internal dan kendaraan dari Excel.","schedule_date":"Tanggal","today":"Hari ini","work_schedule":"Jadwal Kerja Internal","vehicle_schedule":"Jadwal Kendaraan","scheduled_people":"Ada jadwal","people":"orang","no_work_schedule":"Tidak ada jadwal internal pada tanggal ini.","vehicle":"Kendaraan","plate":"Nomor Kendaraan","user":"Pengguna","vehicle_status":"Status","available":"Kosong","in_use_plan":"Jadwal Dipakai","source_file":"Excel sumber","excel_read_error":"Excel tidak dapat dibaca","schedule_refresh":"Perbarui dari Excel","schedule_refreshing":"Memperbarui Excel ke database","schedule_last_sync":"Terakhir diperbarui","schedule_not_synced":"Belum diperbarui","schedule_sync_done":"Jadwal berhasil diperbarui ke database."
    }
}

def ui_language():
    if "user_id" not in session:
        return "ja"
    try:
        conn=db(); row=conn.execute("SELECT language FROM users WHERE id=?",(session["user_id"],)).fetchone(); conn.close()
        return row["language"] if row and row["language"] in ("ja","id") else "ja"
    except Exception:
        return "ja"

def tr(key):
    lang=ui_language()
    return UI_TEXT.get(lang,UI_TEXT["ja"]).get(key,UI_TEXT["ja"].get(key,key))


def _excel_date(value):
    """Convert Excel/COM date values to a Python date when possible."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
        except Exception:
            return None
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except Exception:
            pass
    return None


def _clean_excel_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_xlsx_matrix(path, sheet_name=None, max_rows=300, max_cols=500):
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        rows = min(ws.max_row or 1, max_rows)
        cols = min(ws.max_column or 1, max_cols)
        data = []
        for r in range(1, rows + 1):
            data.append([ws.cell(r, c).value for c in range(1, cols + 1)])
        return data
    finally:
        wb.close()


def _read_xls_via_excel(path, sheet_name, max_rows=300, max_cols=500):
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        raise RuntimeError(".xlsの読込にはMicrosoft Excelとpywin32が必要です。")
    pythoncom.CoInitialize()
    excel = workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(os.path.abspath(path), UpdateLinks=0, ReadOnly=True)
        ws = workbook.Worksheets(sheet_name)
        used = ws.UsedRange
        rows = min(int(used.Rows.Count), max_rows)
        cols = min(int(used.Columns.Count), max_cols)
        values = ws.Range(ws.Cells(1, 1), ws.Cells(rows, cols)).Value
        if rows == 1:
            values = (values,)
        return [list(row) for row in values]
    finally:
        if workbook is not None:
            try: workbook.Close(False)
            except Exception: pass
        if excel is not None:
            try: excel.Quit()
            except Exception: pass
        pythoncom.CoUninitialize()


def _read_schedule_matrix(path, sheet_name=None):
    if not os.path.exists(path):
        raise FileNotFoundError(os.path.basename(path) + " が見つかりません。")
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx_matrix(path, sheet_name)
    if ext == ".xls":
        return _read_xls_via_excel(path, sheet_name or "社内")
    raise ValueError("対応していないExcel形式です。")


def _find_date_header(matrix, target_date, row_limit=20):
    best_row = None
    best_count = 0
    for r, row in enumerate(matrix[:row_limit]):
        count = sum(1 for v in row if _excel_date(v))
        if count > best_count:
            best_row, best_count = r, count
    if best_row is None or best_count == 0:
        return None, None
    for c, v in enumerate(matrix[best_row]):
        if _excel_date(v) == target_date:
            return best_row, c
    return best_row, None


def read_internal_work_schedule(target_date):
    matrix = _read_schedule_matrix(WORK_SCHEDULE_FILE, "社内")
    header_row, target_col = _find_date_header(matrix, target_date)
    if target_col is None:
        return []
    excluded = {"工事協力会社", "日程未確定"}
    results = []
    current_name = None
    current_items = []

    def flush():
        if current_name and current_name not in excluded and current_items:
            unique = []
            for item in current_items:
                if item and item not in unique:
                    unique.append(item)
            if unique:
                results.append({"name": current_name, "items": unique})

    for row in matrix[(header_row or 0) + 1:]:
        first = _clean_excel_text(row[0] if row else None)
        if first:
            flush()
            current_name = first
            current_items = []
            if current_name == "日程未確定":
                break
        if current_name and current_name not in excluded and target_col < len(row):
            item = _clean_excel_text(row[target_col])
            if item:
                current_items.append(item)
    flush()
    return results


def read_vehicle_schedule(target_date):
    matrix = _read_schedule_matrix(VEHICLE_SCHEDULE_FILE)
    header_row, target_col = _find_date_header(matrix, target_date, row_limit=5)
    if target_col is None:
        return []
    vehicles = []
    # The vehicle workbook uses repeating 3-row groups: vehicle / plate / assigned owner.
    start = (header_row or 0) + 2
    r = start
    while r < len(matrix):
        row = matrix[r]
        vehicle = _clean_excel_text(row[0] if row else None)
        if not vehicle:
            r += 3
            continue
        plate = _clean_excel_text(matrix[r + 1][0] if r + 1 < len(matrix) and matrix[r + 1] else None)
        default_user = _clean_excel_text(matrix[r + 2][0] if r + 2 < len(matrix) and matrix[r + 2] else None)
        assigned = _clean_excel_text(row[target_col] if target_col < len(row) else None)
        vehicles.append({
            "vehicle": vehicle,
            "plate": plate,
            "default_user": default_user,
            "assigned": assigned,
            "available": not bool(assigned),
        })
        r += 3
    return vehicles

@app.context_processor
def inject_ui_text():
    return {"tr":tr,"ui_lang":ui_language()}


def current_user():
    if "user_id" not in session: return None
    conn = db(); row = conn.execute("SELECT * FROM users WHERE id=?", (session["user_id"],)).fetchone(); conn.close(); return row


@app.before_request
def ensure_database():
    init_db()
    if request.endpoint not in ("setup", "static") and not has_users():
        return redirect(url_for("setup"))


@app.route("/setup", methods=["GET", "POST"])
def setup():
    if has_users(): return redirect(url_for("login"))
    if request.method == "POST":
        name=request.form.get("name","").strip(); employee_no=request.form.get("employee_no","").strip(); username=request.form.get("username","").strip(); password=request.form.get("password","")
        if not name or not username or len(password) < 8:
            flash("氏名、ユーザー名、8文字以上のパスワードを入力してください。", "error")
        else:
            conn=db(); conn.execute("INSERT INTO users(employee_no,name,username,password_hash,role) VALUES(?,?,?,?,?)", (employee_no or None,name,username,generate_password_hash(password),"admin")); conn.commit(); conn.close()
            flash("管理者アカウントを作成しました。", "success"); return redirect(url_for("login"))
    return render_template("setup.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username=request.form.get("username","").strip(); password=request.form.get("password","")
        conn=db(); user=conn.execute("SELECT * FROM users WHERE username=? AND active=1", (username,)).fetchone(); conn.close()
        if user and check_password_hash(user["password_hash"], password):
            session.clear(); session.update(user_id=user["id"], name=user["name"], role=user["role"]); return redirect(url_for("dashboard"))
        flash("ユーザー名またはパスワードが正しくありません。", "error")
    return render_template("login.html")


@app.route("/logout")
def logout(): session.clear(); return redirect(url_for("login"))


def import_schedule_to_database():
    """Read both schedule Excel files once and replace the SQLite schedule cache."""
    work_matrix = _read_schedule_matrix(WORK_SCHEDULE_FILE, "社内")
    vehicle_matrix = _read_schedule_matrix(VEHICLE_SCHEDULE_FILE)

    work_header, _ = _find_date_header(work_matrix, date.today())
    if work_header is None:
        # Find the row with the most date cells even when today's date is outside the workbook.
        work_header, _ = _find_date_header(work_matrix, date(1900, 1, 1))
    vehicle_header, _ = _find_date_header(vehicle_matrix, date.today(), row_limit=5)
    if vehicle_header is None:
        vehicle_header, _ = _find_date_header(vehicle_matrix, date(1900, 1, 1), row_limit=5)
    if work_header is None or vehicle_header is None:
        raise ValueError("Excelの日付行を確認できません。")

    work_records = []
    excluded = {"工事協力会社", "日程未確定"}
    for c, raw_date in enumerate(work_matrix[work_header]):
        d = _excel_date(raw_date)
        if not d:
            continue
        current_name = None
        person_order = 0
        for row in work_matrix[work_header + 1:]:
            first = _clean_excel_text(row[0] if row else None)
            if first:
                current_name = first
                person_order += 1
                if current_name == "日程未確定":
                    break
            if current_name and current_name not in excluded and c < len(row):
                item = _clean_excel_text(row[c])
                if item:
                    work_records.append((d.isoformat(), current_name, item, person_order))

    vehicle_records = []
    date_columns = [(c, _excel_date(v)) for c, v in enumerate(vehicle_matrix[vehicle_header])]
    date_columns = [(c, d) for c, d in date_columns if d]
    r = vehicle_header + 2
    vehicle_order = 0
    while r < len(vehicle_matrix):
        row = vehicle_matrix[r]
        vehicle = _clean_excel_text(row[0] if row else None)
        if vehicle:
            vehicle_order += 1
            plate = _clean_excel_text(vehicle_matrix[r + 1][0] if r + 1 < len(vehicle_matrix) and vehicle_matrix[r + 1] else None)
            default_user = _clean_excel_text(vehicle_matrix[r + 2][0] if r + 2 < len(vehicle_matrix) and vehicle_matrix[r + 2] else None)
            for c, d in date_columns:
                assigned = _clean_excel_text(row[c] if c < len(row) else None)
                vehicle_records.append((d.isoformat(), vehicle, plate, default_user, assigned, vehicle_order))
        r += 3

    conn = db()
    try:
        conn.execute("BEGIN")
        conn.execute("DELETE FROM schedule_work_cache")
        conn.execute("DELETE FROM schedule_vehicle_cache")
        conn.executemany("INSERT INTO schedule_work_cache(schedule_date,person_name,item_text,sort_order) VALUES(?,?,?,?)", work_records)
        conn.executemany("INSERT INTO schedule_vehicle_cache(schedule_date,vehicle,plate,default_user,assigned,sort_order) VALUES(?,?,?,?,?,?)", vehicle_records)
        conn.execute("INSERT OR REPLACE INTO schedule_sync_state(id,synced_at,work_file_mtime,vehicle_file_mtime) VALUES(1,?,?,?)", (
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            os.path.getmtime(WORK_SCHEDULE_FILE), os.path.getmtime(VEHICLE_SCHEDULE_FILE)))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def read_cached_work_schedule(target_date):
    conn = db()
    rows = conn.execute("SELECT person_name,item_text,sort_order FROM schedule_work_cache WHERE schedule_date=? ORDER BY sort_order,id", (target_date.isoformat(),)).fetchall()
    conn.close()
    grouped = []
    by_name = {}
    for row in rows:
        name = row["person_name"]
        if name not in by_name:
            rec = {"name": name, "items": []}
            by_name[name] = rec
            grouped.append(rec)
        if row["item_text"] not in by_name[name]["items"]:
            by_name[name]["items"].append(row["item_text"])
    return grouped


def read_cached_vehicle_schedule(target_date):
    conn = db()
    rows = conn.execute("SELECT vehicle,plate,default_user,assigned FROM schedule_vehicle_cache WHERE schedule_date=? ORDER BY sort_order,id", (target_date.isoformat(),)).fetchall()
    conn.close()
    return [{"vehicle":r["vehicle"], "plate":r["plate"] or "", "default_user":r["default_user"] or "", "assigned":r["assigned"] or "", "available":not bool(r["assigned"])} for r in rows]


@app.post("/schedule/import")
@login_required
def schedule_import():
    try:
        import_schedule_to_database()
        flash(tr("schedule_sync_done"), "success")
    except Exception as e:
        flash(f"{tr('excel_read_error')}：{e}", "error")
    return redirect(url_for("schedule_view", date=request.form.get("date") or date.today().isoformat()))


@app.route("/schedule")
@login_required
def schedule_view():
    raw = (request.args.get("date") or date.today().isoformat()).strip()
    try:
        selected = datetime.strptime(raw, "%Y-%m-%d").date()
    except ValueError:
        selected = date.today()
    work_error = vehicle_error = None
    work_rows = read_cached_work_schedule(selected)
    vehicle_rows = read_cached_vehicle_schedule(selected)
    conn = db()
    sync_row = conn.execute("SELECT synced_at FROM schedule_sync_state WHERE id=1").fetchone()
    conn.close()
    last_sync = sync_row["synced_at"] if sync_row and sync_row["synced_at"] else None
    return render_template(
        "schedule.html",
        selected_date=selected.isoformat(),
        prev_date=(selected - timedelta(days=1)).isoformat(),
        next_date=(selected + timedelta(days=1)).isoformat(),
        work_rows=work_rows,
        vehicle_rows=vehicle_rows,
        work_error=work_error,
        vehicle_error=vehicle_error,
        work_file=os.path.basename(WORK_SCHEDULE_FILE),
        vehicle_file=os.path.basename(VEHICLE_SCHEDULE_FILE),
        last_sync=last_sync,
    )


@app.route("/")
@login_required
def dashboard():
    month=request.args.get("month") or date.today().strftime("%Y-%m")
    conn=db()
    reports=conn.execute('''SELECT d.*, u.name,
        COALESCE(SUM(e.work_minutes),0) AS total_minutes,
        COUNT(e.id) AS entry_count
        FROM daily_reports d JOIN users u ON u.id=d.user_id
        LEFT JOIN work_entries e ON e.daily_report_id=d.id
        WHERE d.user_id=? AND substr(d.work_date,1,7)=?
        GROUP BY d.id ORDER BY d.work_date DESC''', (session["user_id"],month)).fetchall()
    summary=conn.execute('''SELECT COUNT(DISTINCT d.id) days, COALESCE(SUM(e.work_minutes),0) work_minutes, COUNT(e.id) entries
        FROM daily_reports d LEFT JOIN work_entries e ON e.daily_report_id=d.id
        WHERE d.user_id=? AND substr(d.work_date,1,7)=?''', (session["user_id"],month)).fetchone()

    # Manual time classifications are shown separately from normal work time.
    cats = ("み", "残", "早", "F", "深")
    month_adjust = {c: 0.0 for c in cats}
    for row in conn.execute('''SELECT a.category, COALESCE(SUM(a.hours),0) AS hours
        FROM time_adjustments a JOIN daily_reports d ON d.id=a.daily_report_id
        WHERE d.user_id=? AND substr(d.work_date,1,7)=? AND a.category <> '-'
        GROUP BY a.category''', (session["user_id"],month)).fetchall():
        if row["category"] in month_adjust:
            month_adjust[row["category"]] = float(row["hours"] or 0)

    report_adjust = {}
    report_span = {}
    for r in reports:
        rid = r["id"]
        report_adjust[rid] = {c: 0.0 for c in cats}
        for row in conn.execute("SELECT category, COALESCE(SUM(hours),0) hours FROM time_adjustments WHERE daily_report_id=? AND category <> '-' GROUP BY category", (rid,)).fetchall():
            if row["category"] in report_adjust[rid]:
                report_adjust[rid][row["category"]] = float(row["hours"] or 0)
        span = conn.execute("SELECT MIN(start_time) first_start, MAX(end_time) last_end FROM work_entries WHERE daily_report_id=?", (rid,)).fetchone()
        report_span[rid] = {"start": span["first_start"] if span else None, "end": span["last_end"] if span else None}

    conn.close()
    return render_template("dashboard.html", reports=reports, summary=summary, month=month, user=current_user(), month_adjust=month_adjust, report_adjust=report_adjust, report_span=report_span)


def get_jobs():
    conn=db(); rows=conn.execute("SELECT * FROM jobs WHERE active=1 ORDER BY job_no").fetchall(); conn.close(); return rows


@app.route("/day/new", methods=["GET", "POST"])
@login_required
def new_day():
    if request.method == "POST": return save_day(None)
    return render_template("day_form.html", report=None, entries=[], adjustments=[], jobs=get_jobs(), work_codes=WORK_CODES, today=date.today().isoformat())


@app.route("/day/<int:rid>/edit", methods=["GET", "POST"])
@login_required
def edit_day(rid):
    conn=db(); report=conn.execute("SELECT * FROM daily_reports WHERE id=?",(rid,)).fetchone()
    if not report or (report["user_id"] != session["user_id"] and session.get("role") != "admin"):
        conn.close(); flash("日報が見つかりません。","error"); return redirect(url_for("dashboard"))
    entries=conn.execute("SELECT * FROM work_entries WHERE daily_report_id=? ORDER BY sort_order,id",(rid,)).fetchall(); adjustments=conn.execute("SELECT * FROM time_adjustments WHERE daily_report_id=? ORDER BY sort_order,id",(rid,)).fetchall(); conn.close()
    if request.method == "POST": return save_day(rid)
    return render_template("day_form.html", report=report, entries=entries, adjustments=adjustments, jobs=get_jobs(), work_codes=WORK_CODES, today=date.today().isoformat())


def save_day(rid):
    work_date=request.form.get("work_date",""); attendance_start=request.form.get("attendance_start",""); attendance_end=request.form.get("attendance_end",""); remarks=request.form.get("remarks","").strip()
    job_ids=request.form.getlist("job_id[]"); work_codes=request.form.getlist("work_code[]"); contents=request.form.getlist("content[]"); locations=request.form.getlist("location[]"); starts=request.form.getlist("start_time[]"); ends=request.form.getlist("end_time[]"); adjust_types=request.form.getlist("adjust_type[]"); adjust_hours=request.form.getlist("adjust_hours[]"); adjust_entry_orders=request.form.getlist("adjust_entry_order[]")
    if not work_date or not attendance_start or not attendance_end:
        flash("日付と勤務時間を入力してください。","error"); return redirect(request.url)
    if not valid_quarter(attendance_start) or not valid_quarter(attendance_end):
        flash("時刻は15分単位で入力してください。","error"); return redirect(request.url)
    rows=[]; intervals=[]
    count=max(len(job_ids),len(work_codes),len(contents),len(locations),len(starts),len(ends))
    for i in range(count):
        job_id=(job_ids[i] if i<len(job_ids) else "").strip(); code=(work_codes[i] if i<len(work_codes) else "").strip(); content=(contents[i] if i<len(contents) else "").strip(); location=(locations[i] if i<len(locations) else "").strip(); start=(starts[i] if i<len(starts) else "").strip(); end=(ends[i] if i<len(ends) else "").strip()
        if not any([job_id,code,content,location,start,end]): continue
        if code and code not in WORK_CODES:
            flash(f"{i+1}行目：業務コードが正しくありません。","error"); return redirect(request.url)
        if not content:
            flash(f"{i+1}行目：内容を入力してください。","error"); return redirect(request.url)
        if code not in ("H1","H2") and (not start or not end):
            flash(f"{i+1}行目：開始・終了時刻を入力してください。","error"); return redirect(request.url)
        if not valid_quarter(start) or not valid_quarter(end):
            flash(f"{i+1}行目：時刻は15分単位で入力してください。","error"); return redirect(request.url)
        if code == "H1": start,end="08:30","12:00"
        if code == "H2": start,end="13:00","17:30"
        iv=interval(start,end)
        if iv:
            for prev_i,prev in intervals:
                if intervals_overlap(iv,prev):
                    flash(f"{prev_i+1}行目と{i+1}行目の時間が重複しています。","error"); return redirect(request.url)
            intervals.append((i,iv))
        rows.append((job_id or None,code,content,location,start,end,minutes_between(start,end),i))
    if not rows:
        flash("作業を1件以上入力してください。","error"); return redirect(request.url)

    conn=db()
    # unique day per employee
    existing_day=conn.execute("SELECT id FROM daily_reports WHERE user_id=? AND work_date=?",(session["user_id"],work_date)).fetchone()
    if rid is None and existing_day:
        conn.close(); flash("この日の日報は既に登録されています。編集画面を開きます。","error"); return redirect(url_for("edit_day",rid=existing_day["id"]))
    if rid is None:
        conn.execute("INSERT INTO daily_reports(user_id,work_date,attendance_start,attendance_end,remarks) VALUES(?,?,?,?,?)",(session["user_id"],work_date,attendance_start,attendance_end,remarks)); rid=conn.execute("SELECT last_insert_rowid() id").fetchone()["id"]
    else:
        owner=conn.execute("SELECT user_id FROM daily_reports WHERE id=?",(rid,)).fetchone()
        if not owner or (owner["user_id"] != session["user_id"] and session.get("role") != "admin"):
            conn.close(); flash("この操作は許可されていません。","error"); return redirect(url_for("dashboard"))
        conn.execute("UPDATE daily_reports SET work_date=?,attendance_start=?,attendance_end=?,remarks=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(work_date,attendance_start,attendance_end,remarks,rid)); conn.execute("DELETE FROM work_entries WHERE daily_report_id=?",(rid,)); conn.execute("DELETE FROM time_adjustments WHERE daily_report_id=?",(rid,))
    for job_id,code,content,location,start,end,mins,order in rows:
        snapshot=None
        if job_id:
            job=conn.execute("SELECT job_no FROM jobs WHERE id=?",(job_id,)).fetchone()
            if not job: conn.rollback(); conn.close(); flash("JOB番号がデータベースに見つかりません。","error"); return redirect(request.url)
            snapshot=job["job_no"]
        conn.execute("INSERT INTO work_entries(daily_report_id,job_id,job_no_snapshot,work_code,content,location,start_time,end_time,work_minutes,sort_order) VALUES(?,?,?,?,?,?,?,?,?,?)",(rid,job_id,snapshot,code,content,location,start,end,mins,order))
    allowed_adjust={"-","残","み","早","F","深"}
    work_day = datetime.strptime(work_date, "%Y-%m-%d").date()
    has_travel = any(code == "E" for _,code,_,_,_,_,_,_ in rows)
    for i, typ in enumerate(adjust_types):
        typ=(typ or "-").strip()
        try:
            entry_order=int(adjust_entry_orders[i]) if i < len(adjust_entry_orders) else 0
        except (ValueError, TypeError):
            entry_order=0
        hv=(adjust_hours[i] if i<len(adjust_hours) else "").strip()
        if typ not in allowed_adjust:
            conn.rollback(); conn.close(); flash(f"時間区分{i+1}：区分が正しくありません。","error"); return redirect(request.url)
        if typ == "-":
            hours=0.0
        else:
            try:
                hours=float(hv)
            except ValueError:
                conn.rollback(); conn.close(); flash(f"時間区分{i+1}：時間を数値で入力してください。","error"); return redirect(request.url)
            if hours <= 0:
                conn.rollback(); conn.close(); flash(f"時間区分{i+1}：時間を入力してください。","error"); return redirect(request.url)
            # Company rule: all self-reported time is in 15-minute increments.
            if abs(hours * 4 - round(hours * 4)) > 1e-9:
                conn.rollback(); conn.close(); flash(f"時間区分{i+1}：時間は0.25時間（15分）単位で入力してください。","error"); return redirect(request.url)
            if typ == "み" and hours > 1.0:
                conn.rollback(); conn.close(); flash("みなし時間は1日最大1.00時間です。","error"); return redirect(request.url)
            if typ == "み" and work_day.weekday() >= 5:
                conn.rollback(); conn.close(); flash("土日休業日には『み』は入力できません。","error"); return redirect(request.url)
            if typ == "F":
                row_code = next((r[1] for r in rows if r[7] == entry_order), "")
                if row_code != "E":
                    conn.rollback(); conn.close(); flash(f"作業{entry_order+1}：F時間外移動を入力する場合、この作業の業務コードをE（移動）にしてください。","error"); return redirect(request.url)
        conn.execute("INSERT INTO time_adjustments(daily_report_id,category,hours,sort_order,work_entry_order) VALUES(?,?,?,?,?)",(rid,typ,hours,i,entry_order))
    conn.execute("UPDATE daily_reports SET sync_status='unsynced' WHERE id=?",(rid,))
    conn.commit(); conn.close()
    # 保存ボタンはデータベース保存のみ。会社Excelへの書き込みは別ボタンで実行する。
    flash("業務日報を保存しました。Excelにはまだ反映していません。", "success")
    return redirect(url_for("dashboard",month=work_date[:7]))


@app.post("/day/<int:rid>/excel-sync")
@login_required
def excel_sync_day(rid):
    conn = db()
    report = conn.execute("SELECT * FROM daily_reports WHERE id=?", (rid,)).fetchone()
    conn.close()
    if not report:
        flash("日報データが見つかりません。", "error")
        return redirect(url_for("dashboard"))
    if report["user_id"] != session["user_id"] and session.get("role") != "admin":
        flash("この操作は許可されていません。", "error")
        return redirect(url_for("dashboard"))
    try:
        excel_path = sync_report_to_company_excel(rid)
        flash(f"Excelへ保存しました：{os.path.basename(excel_path)}", "success")
    except Exception as e:
        flash(f"Excelへの保存に失敗しました：{e}", "error")
    return redirect(url_for("dashboard", month=report["work_date"][:7]))


@app.post("/excel-sync-all")
@login_required
def excel_sync_all():
    month = (request.form.get("month") or "").strip()
    if not month or len(month) != 7:
        month = date.today().strftime("%Y-%m")
    conn = db()
    if session.get("role") == "admin" and request.form.get("all_users") == "1":
        reports = conn.execute("SELECT id, work_date FROM daily_reports WHERE substr(work_date,1,7)=? ORDER BY work_date,id", (month,)).fetchall()
    else:
        reports = conn.execute("SELECT id, work_date FROM daily_reports WHERE user_id=? AND substr(work_date,1,7)=? ORDER BY work_date,id", (session["user_id"], month)).fetchall()
    conn.close()
    if not reports:
        flash(f"{month} の日報がありません。", "error")
        return redirect(url_for("dashboard", month=month))
    success = 0
    errors = []
    excel_name = ""
    for r in reports:
        try:
            excel_path = sync_report_to_company_excel(r["id"])
            excel_name = os.path.basename(excel_path)
            success += 1
        except Exception as e:
            errors.append(f"{r['work_date']}: {e}")
    if success:
        flash(f"{month} の日報 {success}日分をExcelへ一括保存しました：{excel_name}", "success")
    if errors:
        preview = " / ".join(errors[:3])
        if len(errors) > 3:
            preview += f" / 他{len(errors)-3}件"
        flash(f"Excelへ保存できなかった日報があります：{preview}", "error")
    return redirect(url_for("dashboard", month=month))


@app.post("/day/<int:rid>/delete")
@login_required
def delete_day(rid):
    conn=db(); r=conn.execute("SELECT * FROM daily_reports WHERE id=?",(rid,)).fetchone()
    if r and (r["user_id"]==session["user_id"] or session.get("role")=="admin"):
        conn.execute("DELETE FROM daily_reports WHERE id=?",(rid,)); conn.commit(); flash("日報を削除しました。","success")
    conn.close(); return redirect(url_for("dashboard"))


@app.route("/jobs", methods=["GET", "POST"])
@admin_required
def jobs():
    conn=db()
    if request.method == "POST":
        job_no=request.form.get("job_no","").strip(); project=request.form.get("project_name","").strip(); client=request.form.get("client_name","").strip(); default_content=request.form.get("default_content","").strip()
        if not job_no: flash("JOB番号を入力してください。","error")
        else:
            try:
                conn.execute("INSERT INTO jobs(job_no,project_name,client_name,default_content) VALUES(?,?,?,?)",(job_no,project,client,default_content)); conn.commit(); flash("JOB番号を登録しました。","success")
            except sqlite3.IntegrityError: flash("このJOB番号は既に登録されています。","error")
    rows=conn.execute("SELECT * FROM jobs ORDER BY active DESC, job_no").fetchall(); conn.close(); return render_template("jobs.html",jobs=rows)


@app.post("/jobs/<int:jid>/toggle")
@admin_required
def toggle_job(jid):
    conn=db(); conn.execute("UPDATE jobs SET active=CASE active WHEN 1 THEN 0 ELSE 1 END WHERE id=?",(jid,)); conn.commit(); conn.close(); return redirect(url_for("jobs"))


@app.route("/users", methods=["GET", "POST"])
@admin_required
def users():
    if request.method == "POST":
        name=request.form.get("name","").strip(); employee_no=request.form.get("employee_no","").strip(); username=request.form.get("username","").strip(); password=request.form.get("password",""); role=request.form.get("role","employee")
        if not employee_no or not name or not username or len(password)<8:
            flash("社員番号、氏名、ユーザー名、8文字以上のパスワードを入力してください。","error")
        else:
            try:
                conn=db(); cur=conn.execute("INSERT INTO users(employee_no,name,username,password_hash,role) VALUES(?,?,?,?,?)",(employee_no,name,username,generate_password_hash(password),role if role in ("admin","employee") else "employee")); conn.commit(); uid=cur.lastrowid; conn.close()
                try:
                    excel_path, created = create_employee_excel(employee_no, name)
                    if created:
                        flash(f"アカウントと社員番号 {employee_no} のExcelを作成しました：{os.path.basename(excel_path)}","success")
                    else:
                        flash("アカウントを作成しました。対応するExcelは既にあります。","success")
                except Exception as e:
                    flash(f"アカウントは作成しましたが、Excelを作成できませんでした：{e}","error")
            except sqlite3.IntegrityError:
                try: conn.close()
                except Exception: pass
                flash("ユーザー名または社員番号は既に使用されています。","error")
    conn=db(); rows=conn.execute("SELECT * FROM users ORDER BY name").fetchall(); conn.close(); return render_template("users.html",users=rows)


@app.post("/users/excel-create-missing")
@admin_required
def create_missing_user_excels():
    conn=db(); rows=conn.execute("SELECT employee_no,name FROM users WHERE employee_no IS NOT NULL AND TRIM(employee_no)<>'' ORDER BY name").fetchall(); conn.close()
    created=0; skipped=0; errors=[]
    for u in rows:
        try:
            path, was_created = create_employee_excel(u["employee_no"], u["name"])
            if was_created: created += 1
            else: skipped += 1
        except Exception as e:
            errors.append(f"{u['employee_no']} {u['name']}: {e}")
    if created:
        flash(f"不足していた業務日報Excelを {created} 件作成しました。既存 {skipped} 件は変更していません。","success")
    elif not errors:
        flash(f"全員のExcelは既にあります（{skipped}件）。","success")
    if errors:
        flash("Excelを作成できないアカウントがあります：" + " / ".join(errors), "error")
    return redirect(url_for("users"))


@app.route("/settings", methods=["GET", "POST"])
@login_required
def settings():
    if request.method == "POST":
        theme=request.form.get("theme","light")
        language=request.form.get("language","ja")
        if theme not in ("light","dark"): theme="light"
        if language not in ("ja","id"): language="ja"
        conn=db(); conn.execute("UPDATE users SET theme=?,language=? WHERE id=?",(theme,language,session["user_id"])); conn.commit(); conn.close()
        flash("Pengaturan disimpan." if language=="id" else "設定を保存しました。","success"); return redirect(url_for("settings"))
    return render_template("settings.html", user=current_user())

@app.route("/users/<int:uid>/edit", methods=["GET","POST"])
@admin_required
def edit_user(uid):
    conn=db(); u=conn.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone()
    if not u: conn.close(); flash("アカウントが見つかりません。","error"); return redirect(url_for("users"))
    if request.method=="POST":
        employee_no=request.form.get("employee_no","").strip() or None; name=request.form.get("name","").strip(); username=request.form.get("username","").strip(); role=request.form.get("role","employee"); password=request.form.get("password","")
        try:
            if password:
                if len(password)<8: raise ValueError("パスワードは8文字以上にしてください。")
                conn.execute("UPDATE users SET employee_no=?,name=?,username=?,role=?,password_hash=? WHERE id=?",(employee_no,name,username,role,generate_password_hash(password),uid))
            else: conn.execute("UPDATE users SET employee_no=?,name=?,username=?,role=? WHERE id=?",(employee_no,name,username,role,uid))
            conn.commit(); conn.close(); flash("アカウントを更新しました。","success"); return redirect(url_for("users"))
        except (sqlite3.IntegrityError,ValueError) as e: flash(str(e) if isinstance(e,ValueError) else "ユーザー名または社員番号は既に使用されています。","error")
    conn.close(); return render_template("user_edit.html",u=u)

@app.route("/jobs/<int:jid>/edit", methods=["GET","POST"])
@admin_required
def edit_job(jid):
    conn=db(); j=conn.execute("SELECT * FROM jobs WHERE id=?",(jid,)).fetchone()
    if not j: conn.close(); flash("JOBが見つかりません。","error"); return redirect(url_for("jobs"))
    if request.method=="POST":
        try:
            conn.execute("UPDATE jobs SET job_no=?,project_name=?,client_name=?,default_content=?,active=? WHERE id=?",(request.form.get("job_no","").strip(),request.form.get("project_name","").strip(),request.form.get("client_name","").strip(),request.form.get("default_content","").strip(),1 if request.form.get("active")=="1" else 0,jid)); conn.commit(); conn.close(); flash("JOBを更新しました。","success"); return redirect(url_for("jobs"))
        except sqlite3.IntegrityError: flash("このJOB番号は既に登録されています。","error")
    conn.close(); return render_template("job_edit.html",j=j)

@app.post("/jobs/<int:jid>/delete")
@admin_required
def delete_job(jid):
    conn=db()
    job=conn.execute("SELECT job_no FROM jobs WHERE id=?",(jid,)).fetchone()
    if not job:
        conn.close(); flash("JOBが見つかりません。","error"); return redirect(url_for("jobs"))
    # Keep the JOB number snapshot in historical work entries, but detach the
    # foreign-key reference so the master JOB can be deleted completely.
    conn.execute("UPDATE work_entries SET job_no_snapshot=COALESCE(NULLIF(job_no_snapshot,''), ?), job_id=NULL WHERE job_id=?",(job["job_no"],jid))
    conn.execute("DELETE FROM jobs WHERE id=?",(jid,))
    conn.commit(); conn.close()
    flash("JOBを削除しました。過去の日報はJOB番号の記録を保持します。","success")
    return redirect(url_for("jobs"))

@app.post("/day/<int:rid>/copy")
@login_required
def copy_day(rid):
    target=request.form.get("target_date","")
    try: datetime.strptime(target,"%Y-%m-%d")
    except: flash("コピー先の日付が正しくありません。","error"); return redirect(url_for("dashboard"))
    conn=db(); r=conn.execute("SELECT * FROM daily_reports WHERE id=? AND user_id=?",(rid,session["user_id"])).fetchone()
    if not r: conn.close(); flash("日報が見つかりません。","error"); return redirect(url_for("dashboard"))
    if conn.execute("SELECT 1 FROM daily_reports WHERE user_id=? AND work_date=?",(session["user_id"],target)).fetchone(): conn.close(); flash("コピー先の日付には既に日報があります。","error"); return redirect(url_for("dashboard",month=target[:7]))
    cur=conn.execute("INSERT INTO daily_reports(user_id,work_date,attendance_start,attendance_end,remarks,sync_status) VALUES(?,?,?,?,?,\'unsynced\')",(session["user_id"],target,r["attendance_start"],r["attendance_end"],r["remarks"])); nid=cur.lastrowid
    for e in conn.execute("SELECT * FROM work_entries WHERE daily_report_id=? ORDER BY sort_order",(rid,)).fetchall(): conn.execute("INSERT INTO work_entries(daily_report_id,job_id,job_no_snapshot,work_code,content,location,start_time,end_time,work_minutes,sort_order) VALUES(?,?,?,?,?,?,?,?,?,?)",(nid,e["job_id"],e["job_no_snapshot"],e["work_code"],e["content"],e["location"],e["start_time"],e["end_time"],e["work_minutes"],e["sort_order"]))
    for x in conn.execute("SELECT * FROM time_adjustments WHERE daily_report_id=?",(rid,)).fetchall(): conn.execute("INSERT INTO time_adjustments(daily_report_id,category,hours,sort_order,work_entry_order) VALUES(?,?,?,?,?)",(nid,x["category"],x["hours"],x["sort_order"],x["work_entry_order"]))
    conn.commit(); conn.close(); flash("日報をコピーしました。","success"); return redirect(url_for("edit_day",rid=nid))

@app.get("/backup-db")
@admin_required
def backup_db():
    if not os.path.exists(DB_PATH): flash("データベースがありません。","error"); return redirect(url_for("settings"))
    name=f"daily_report_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
    return send_file(DB_PATH,as_attachment=True,download_name=name)


def _excel_hhmm(v):
    if v is None or v=="": return ""
    if isinstance(v,datetime): return v.strftime("%H:%M")
    if isinstance(v,(int,float)):
        mins=round((float(v)%1)*1440); return f"{(mins//60)%24:02d}:{mins%60:02d}"
    text=str(v).strip()
    try: return datetime.strptime(text[:5],"%H:%M").strftime("%H:%M")
    except: return ""

@app.post("/excel-import")
@login_required
def excel_import():
    month=(request.form.get("month") or date.today().strftime("%Y-%m")).strip()
    try:
        y,m=map(int,month.split("-")); probe=f"{y:04d}-{m:02d}-01"
        conn0=db(); current_user=conn0.execute("SELECT employee_no FROM users WHERE id=?",(session["user_id"],)).fetchone(); conn0.close()
        if not current_user or not str(current_user["employee_no"] or "").strip():
            raise ValueError("アカウントに社員番号が設定されていません。")
        excel_path=find_excel_file(probe, current_user["employee_no"])
    except Exception as e: flash(f"Excel取込に失敗しました：{e}","error"); return redirect(url_for("dashboard",month=month))
    try:
        import pythoncom, win32com.client
        pythoncom.CoInitialize(); xl=win32com.client.DispatchEx("Excel.Application"); xl.Visible=False; xl.DisplayAlerts=False; wb=xl.Workbooks.Open(excel_path,UpdateLinks=0,ReadOnly=True); ws=wb.Worksheets(FULLWIDTH_MONTH_SHEETS[m])
        conn=db(); imported=0
        import calendar
        for day in range(1,calendar.monthrange(y,m)[1]+1):
            base=7+(day-1)*22; ast=_excel_hhmm(ws.Cells(base+1,3).Value); aen=_excel_hhmm(ws.Cells(base+1,5).Value); entries=[]; adjs=[]
            for slot in range(7):
                top=base+5+slot*2; ident=str(ws.Cells(top,2).Value or '').strip(); content=str(ws.Cells(top,3).Value or '').strip(); loc=str(ws.Cells(top,6).Value or '').strip(); st=_excel_hhmm(ws.Cells(top+1,3).Value); en=_excel_hhmm(ws.Cells(top+1,5).Value)
                if any((ident,content,loc,st,en)):
                    job=conn.execute("SELECT * FROM jobs WHERE job_no=?",(ident,)).fetchone(); job_id=job["id"] if job else None; code="" if job else (ident if ident in WORK_CODES else ""); snapshot=ident if job else None; entries.append((job_id,snapshot,code,content,loc,st,en,minutes_between(st,en),slot))
                    for pos in (0,1):
                        typ=str(ws.Cells(top+pos,9).Value or '-').strip(); hv=ws.Cells(top+pos,10).Value
                        if typ in ('み','残','早','F','深') and hv not in (None,''): adjs.append((typ,float(hv),slot*2+pos,slot))
            if not ast and not aen and not entries: continue
            wd=f"{y:04d}-{m:02d}-{day:02d}"; old=conn.execute("SELECT id FROM daily_reports WHERE user_id=? AND work_date=?",(session["user_id"],wd)).fetchone()
            if old: rid=old["id"]; conn.execute("UPDATE daily_reports SET attendance_start=?,attendance_end=?,sync_status='imported',updated_at=CURRENT_TIMESTAMP WHERE id=?",(ast or '08:30',aen or '17:30',rid)); conn.execute("DELETE FROM work_entries WHERE daily_report_id=?",(rid,)); conn.execute("DELETE FROM time_adjustments WHERE daily_report_id=?",(rid,))
            else: rid=conn.execute("INSERT INTO daily_reports(user_id,work_date,attendance_start,attendance_end,sync_status) VALUES(?,?,?,?,\'imported\')",(session["user_id"],wd,ast or '08:30',aen or '17:30')).lastrowid
            for e in entries: conn.execute("INSERT INTO work_entries(daily_report_id,job_id,job_no_snapshot,work_code,content,location,start_time,end_time,work_minutes,sort_order) VALUES(?,?,?,?,?,?,?,?,?,?)",(rid,*e))
            for x in adjs: conn.execute("INSERT INTO time_adjustments(daily_report_id,category,hours,sort_order,work_entry_order) VALUES(?,?,?,?,?)",(rid,*x))
            imported+=1
        conn.commit(); conn.close(); wb.Close(False); xl.Quit(); pythoncom.CoUninitialize(); flash(f"Excelから{imported}日分を取り込みました。","success")
    except Exception as e:
        try: wb.Close(False); xl.Quit(); pythoncom.CoUninitialize()
        except: pass
        flash(f"Excel取込に失敗しました：{e}","error")
    return redirect(url_for("dashboard",month=month))

@app.route("/export.xlsx")
@login_required
def export_xlsx():
    month=request.args.get("month") or date.today().strftime("%Y-%m")
    conn=db(); rows=conn.execute('''SELECT d.work_date,d.attendance_start,d.attendance_end,d.remarks,u.employee_no,u.name,
        e.job_no_snapshot,e.work_code,e.content,e.location,e.start_time,e.end_time,e.work_minutes,j.project_name,j.client_name
        FROM daily_reports d JOIN users u ON u.id=d.user_id JOIN work_entries e ON e.daily_report_id=d.id LEFT JOIN jobs j ON j.id=e.job_id
        WHERE d.user_id=? AND substr(d.work_date,1,7)=? ORDER BY d.work_date,e.sort_order,e.start_time''',(session["user_id"],month)).fetchall(); conn.close()
    wb=Workbook(); ws=wb.active; ws.title="業務日報"
    ws["A1"]="NISHIKAWA AUTOMATION　業務日報"; ws["A1"].font=Font(bold=True,size=16); ws.merge_cells("A1:O1")
    headers=["日付","社員番号","氏名","勤務開始","勤務終了","JOB番号","業務コード","内容","場所","開始","終了","時間(H)","案件名","客先名","備考"]
    ws.append(headers); fill=PatternFill("solid",fgColor="1F4E78"); white=Font(color="FFFFFF",bold=True)
    for c in ws[2]: c.fill=fill; c.font=white; c.alignment=Alignment(horizontal="center")
    for r in rows: ws.append([r["work_date"],r["employee_no"],r["name"],r["attendance_start"],r["attendance_end"],r["job_no_snapshot"] or "",r["work_code"],r["content"],r["location"] or "",r["start_time"],r["end_time"],r["work_minutes"]/60,r["project_name"] or "",r["client_name"] or "",r["remarks"] or ""])
    widths=[12,12,18,11,11,18,12,34,18,9,9,10,28,24,30]; thin=Side(style="thin",color="D4DCE6")
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.border=Border(bottom=thin); c.alignment=Alignment(vertical="top",wrap_text=True)
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,as_attachment=True,download_name=f"業務日報_{month}.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/health")
def health(): return {"status":"ok","app":"Nishikawa Daily Report","version":"3.9.7"}


if __name__ == "__main__":
    init_db(); app.run(host="0.0.0.0",port=5000,debug=True)
